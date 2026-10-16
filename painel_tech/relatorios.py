import io
from collections import defaultdict
from datetime import datetime, time, timezone
from zoneinfo import ZoneInfo

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from . import db
from .config import get_config
from .domain import (
    INDICADORES,
    label_intervencao,
    label_prioridade,
    label_ritmo,
    label_status,
    meta_indicador,
)
from .view_utils import aplicar_filtros, construir_linhas, ordenar_prioridade

FUSO_PADRAO = "America/Sao_Paulo"

SECOES = ["resumo", "projetos", "sla", "tempo", "eventos"]
POR_PAGINA_OPCOES = [10, 25, 50, 100]
POR_PAGINA_PADRAO = 25

FORMATO_DATA = "dd/mm/yyyy"
FORMATO_DATA_HORA = "dd/mm/yyyy hh:mm"
FORMATO_DURACAO = "[h]:mm:ss"


def _fuso():
    import os

    nome = os.getenv("TZ_RELATORIOS", FUSO_PADRAO)
    try:
        return ZoneInfo(nome)
    except Exception:
        return timezone.utc


def intervalo(de_data, ate_data):
    fuso = _fuso()
    de = datetime.combine(de_data, time.min, tzinfo=fuso)
    ate = datetime.combine(ate_data, time.max, tzinfo=fuso)
    return de, ate


def _overlap_segundos(sessao, de, ate):
    inicio = sessao["inicio"]
    fim = sessao["fim"] or datetime.now(timezone.utc)
    if fim.tzinfo is None:
        fim = fim.replace(tzinfo=timezone.utc)
    if inicio.tzinfo is None:
        inicio = inicio.replace(tzinfo=timezone.utc)
    inicio = max(inicio, de)
    fim = min(fim, ate)
    segundos = (fim - inicio).total_seconds()
    return max(0, int(segundos))


def _max_conclusao(etapas):
    datas = [e.get("concluida_em") for e in etapas if e.get("concluida_em")]
    return max(datas) if datas else None


def _para_data(valor):
    if isinstance(valor, datetime):
        return valor.date()
    return valor


def construir_relatorio(de, ate, args, cfg=None):
    cfg = cfg or get_config()
    ids = set(db.projetos_por_atividade(de, ate))
    if not ids:
        return _relatorio_vazio(de, ate, [])

    linhas = [
        linha
        for linha in construir_linhas(db.projetos_com_etapas(), cfg)
        if linha["projeto"]["id"] in ids
    ]
    args_filtro = {chave: valor for chave, valor in args.items() if chave not in ("de", "ate", "aba")}
    linhas = aplicar_filtros(linhas, args_filtro)
    ids = {linha["projeto"]["id"] for linha in linhas}
    ordenadas = ordenar_prioridade(linhas)

    sessoes = db.sessoes_periodo(de, ate, list(ids)) if ids else []
    etapas = db.etapas_concluidas_periodo(de, ate, list(ids)) if ids else []
    eventos = db.eventos_periodo(de, ate, list(ids)) if ids else []

    tempo_por_projeto = defaultdict(int)
    sessoes_por_projeto = defaultdict(int)
    sessao_linhas = []
    for sessao in sessoes:
        segundos = _overlap_segundos(sessao, de, ate)
        tempo_por_projeto[sessao["projeto_id"]] += segundos
        sessoes_por_projeto[sessao["projeto_id"]] += 1
        sessao_linhas.append(
            {
                "projeto": sessao.get("projeto_nome"),
                "membro": sessao.get("membro_nome") or "—",
                "inicio": max(sessao["inicio"], de),
                "fim": min(sessao["fim"] or datetime.now(timezone.utc), ate),
                "duracao": _segundos_para_duracao(segundos),
            }
        )

    sla_linhas = []
    etapas_por_projeto = defaultdict(list)
    for etapa in etapas:
        etapas_por_projeto[etapa["projeto_id"]].append(etapa)

    for projeto_id, lista in etapas_por_projeto.items():
        lista.sort(key=lambda e: e["concluida_em"])
        anterior = None
        for etapa in lista:
            conclusao = etapa["concluida_em"]
            prevista = _para_data(etapa.get("data_prevista"))
            atraso = (conclusao.date() - prevista).days if prevista else None
            criada = etapa.get("created_at")
            ciclo = None
            if criada:
                ciclo = _segundos_para_duracao((conclusao - criada).total_seconds())
            cadencia = None
            if anterior is not None:
                cadencia = _segundos_para_duracao((conclusao - anterior).total_seconds())
            anterior = conclusao
            sla_linhas.append(
                {
                    "projeto": etapa.get("projeto_nome"),
                    "etapa": etapa.get("nome"),
                    "responsavel": etapa.get("responsavel_nome") or "—",
                    "prevista": prevista,
                    "concluida_em": conclusao,
                    "atraso": atraso,
                    "cadencia": cadencia,
                    "ciclo": ciclo,
                }
            )

    projetos_linhas = []
    for linha in ordenadas:
        projeto = linha["projeto"]
        analise = linha["analise"]
        etapa_concluida = _max_conclusao(linha["etapas"])
        inicio = _para_data(projeto.get("data_inicio"))
        fim = _para_data(projeto.get("data_fim")) or _para_data(etapa_concluida)
        lead = (fim - inicio).days if inicio and fim else None
        etapas_projeto = linha["etapas"]
        projetos_linhas.append(
            {
                "nome": projeto["nome"],
                "status": label_status(projeto["status"])["label"],
                "status_cor": label_status(projeto["status"])["cor"],
                "responsavel": projeto.get("responsavel_nome") or "—",
                "impacto": meta_indicador("impacto", projeto["impacto"])["label"],
                "urgencia": meta_indicador("urgencia", projeto["urgencia"])["label"],
                "risco": meta_indicador("risco", projeto["risco"])["label"],
                "esforco": meta_indicador("esforco", projeto["esforco"])["label"],
                "prioridade": analise["prioridade"]["label"],
                "prioridade_cor": label_prioridade(analise["prioridade"]["nivel"])["cor"],
                "indice": analise["prioridade"]["indice"],
                "ritmo": label_ritmo(analise["saude"]["ritmo"])["label"],
                "intervencao": label_intervencao(analise["intervencao"]["nivel"])["label"],
                "progresso": analise["saude"]["progresso_pct"],
                "concluidas": sum(
                    1 for e in etapas_projeto if e.get("status") == "concluida"
                ),
                "total_etapas": len(etapas_projeto),
                "vencidas": analise["saude"]["etapas_vencidas"],
                "bloqueios": sum(
                    1
                    for e in etapas_projeto
                    if e.get("bloqueada") or e.get("dep_pendente")
                ),
                "inicio": inicio,
                "prev_fim": _para_data(projeto.get("data_prevista_fim")),
                "fim": _para_data(projeto.get("data_fim")),
                "tempo": _segundos_para_duracao(tempo_por_projeto.get(projeto["id"], 0)),
                "sessoes": sessoes_por_projeto.get(projeto["id"], 0),
                "lead_dias": lead,
            }
        )

    tempo_membro = defaultdict(int)
    for sessao in sessoes:
        membro = sessao.get("membro_nome") or "—"
        tempo_membro[membro] += _overlap_segundos(sessao, de, ate)
    tempo_membro_total = sum(tempo_membro.values())
    membros_linhas = [
        {
            "membro": membro,
            "tempo": _segundos_para_duracao(segundos),
            "segundos": segundos,
            "pct": round(segundos / tempo_membro_total * 100)
            if tempo_membro_total
            else 0,
        }
        for membro, segundos in sorted(
            tempo_membro.items(), key=lambda item: item[1], reverse=True
        )
    ]

    eventos_linhas = [
        {
            "quando": evento.get("created_at"),
            "projeto": evento.get("projeto_nome"),
            "tipo": evento.get("tipo"),
            "descricao": evento.get("descricao"),
            "autor": evento.get("autor_nome") or "—",
        }
        for evento in eventos
    ]

    cadencia_valores = [
        linha["cadencia"].total_seconds()
        for linha in sla_linhas
        if linha["cadencia"] is not None
    ]
    atrasos = [linha["atraso"] for linha in sla_linhas if linha["atraso"] is not None]
    tempo_total = sum(tempo_por_projeto.values())

    resumo_grupos = _montar_resumo(
        ordenadas,
        tempo_total,
        sla_linhas,
        cadencia_valores,
        atrasos,
        len(sessoes),
        len(eventos),
    )
    resumo = [
        ("Período", f"{de.date().strftime('%d/%m/%Y')} a {ate.date().strftime('%d/%m/%Y')}")
    ]
    resumo += [
        (item["rotulo"], item["valor"])
        for grupo in resumo_grupos
        for item in grupo["itens"]
    ]

    return {
        "de": de,
        "ate": ate,
        "resumo": resumo,
        "resumo_grupos": resumo_grupos,
        "projetos": projetos_linhas,
        "sla": sla_linhas,
        "sessoes": sessao_linhas,
        "membros": membros_linhas,
        "eventos": eventos_linhas,
        "total_projetos": len(ordenadas),
    }


def _montar_resumo(
    ordenadas, tempo_total, sla_linhas, cadencia_valores, atrasos, total_sessoes, total_eventos
):
    sla_medio = (
        _duracao_str(sum(cadencia_valores) / len(cadencia_valores))
        if cadencia_valores
        else "—"
    )
    atraso_medio = round(sum(atrasos) / len(atrasos), 1) if atrasos else None
    tempo_medio = (
        _duracao_str(tempo_total / len(ordenadas)) if ordenadas else _duracao_str(0)
    )
    destaque_atraso = None
    if atraso_medio is not None:
        destaque_atraso = "ruim" if atraso_medio > 0 else "ok"
    return [
        {
            "titulo": "Volume",
            "icone": "📦",
            "itens": [
                {"rotulo": "Projetos com atividade", "valor": len(ordenadas)},
                {
                    "rotulo": "Projetos abertos",
                    "valor": sum(
                        1 for l in ordenadas if l["projeto"]["status"] != "finalizado"
                    ),
                },
                {
                    "rotulo": "Projetos finalizados",
                    "valor": sum(
                        1 for l in ordenadas if l["projeto"]["status"] == "finalizado"
                    ),
                },
            ],
        },
        {
            "titulo": "Tempo",
            "icone": "⏱️",
            "itens": [
                {"rotulo": "Tempo total registrado", "valor": _duracao_str(tempo_total)},
                {"rotulo": "Tempo médio por projeto", "valor": tempo_medio},
                {"rotulo": "Sessões de tempo", "valor": total_sessoes},
            ],
        },
        {
            "titulo": "SLA e prazo",
            "icone": "🎯",
            "itens": [
                {"rotulo": "Checklists concluídos", "valor": len(sla_linhas)},
                {"rotulo": "SLA médio entre conclusões", "valor": sla_medio},
                {
                    "rotulo": "Atraso médio vs prazo (dias)",
                    "valor": atraso_medio if atraso_medio is not None else "—",
                    "destaque": destaque_atraso,
                },
            ],
        },
        {
            "titulo": "Eventos",
            "icone": "🗂️",
            "itens": [
                {"rotulo": "Eventos registrados", "valor": total_eventos},
            ],
        },
    ]


def paginar(itens, pagina, por_pagina):
    por_pagina = por_pagina if por_pagina and por_pagina > 0 else 25
    total = len(itens)
    total_paginas = max(1, (total + por_pagina - 1) // por_pagina)
    pagina = min(max(1, int(pagina or 1)), total_paginas)
    inicio = (pagina - 1) * por_pagina
    return {
        "itens": itens[inicio : inicio + por_pagina],
        "total": total,
        "pagina": pagina,
        "por_pagina": por_pagina,
        "total_paginas": total_paginas,
        "inicio": inicio + 1 if total else 0,
        "fim": min(inicio + por_pagina, total),
        "tem_anterior": pagina > 1,
        "tem_proxima": pagina < total_paginas,
    }


def preparar_secao(relatorio, aba, pagina, por_pagina):
    if aba in ("projetos", "sla", "tempo", "eventos"):
        colecao = {
            "projetos": relatorio["projetos"],
            "sla": relatorio["sla"],
            "tempo": relatorio["sessoes"],
            "eventos": relatorio["eventos"],
        }[aba]
        pag = paginar(colecao, pagina, por_pagina)
        secao = {"secao": aba, "pag": pag, "itens": pag["itens"]}
        if aba == "tempo":
            secao["membros"] = relatorio["membros"]
        return secao
    return {"secao": "resumo", "pag": None, "itens": [], "resumo": relatorio["resumo"]}


def _relatorio_vazio(de, ate, filtros):
    resumo_grupos = _montar_resumo([], 0, [], [], [], 0, 0)
    resumo = [
        ("Período", f"{de.date().strftime('%d/%m/%Y')} a {ate.date().strftime('%d/%m/%Y')}")
    ]
    resumo += [
        (item["rotulo"], item["valor"])
        for grupo in resumo_grupos
        for item in grupo["itens"]
    ]
    return {
        "de": de,
        "ate": ate,
        "resumo": resumo,
        "resumo_grupos": resumo_grupos,
        "projetos": [],
        "sla": [],
        "sessoes": [],
        "membros": [],
        "eventos": [],
        "total_projetos": 0,
    }


def _segundos_para_duracao(segundos):
    from datetime import timedelta

    if segundos is None:
        return None
    return timedelta(seconds=int(segundos))


def _duracao_str(segundos):
    if segundos is None:
        return "—"
    total = int(segundos)
    horas = total // 3600
    minutos = (total % 3600) // 60
    if horas:
        return f"{horas}h {minutos:02d}min"
    return f"{minutos}min"


COLUNAS_PROJETOS = [
    ("Projeto", "nome", 34, None),
    ("Status", "status", 13, None),
    ("Responsável", "responsavel", 22, None),
    (INDICADORES["impacto"]["label"], "impacto", 11, None),
    (INDICADORES["urgencia"]["label"], "urgencia", 11, None),
    (INDICADORES["risco"]["label"], "risco", 11, None),
    (INDICADORES["esforco"]["label"], "esforco", 13, None),
    ("Prioridade", "prioridade", 12, None),
    ("Índice", "indice", 9, "0.00"),
    ("Ritmo", "ritmo", 13, None),
    ("Intervenção", "intervencao", 22, None),
    ("Progresso (%)", "progresso", 13, "0"),
    ("Concluídas", "concluidas", 11, "0"),
    ("Total etapas", "total_etapas", 12, "0"),
    ("Vencidas", "vencidas", 10, "0"),
    ("Bloqueios", "bloqueios", 10, "0"),
    ("Início", "inicio", 12, FORMATO_DATA),
    ("Prev. fim", "prev_fim", 12, FORMATO_DATA),
    ("Fim", "fim", 12, FORMATO_DATA),
    ("Tempo gasto", "tempo", 14, FORMATO_DURACAO),
    ("Sessões", "sessoes", 10, "0"),
    ("Lead time (dias)", "lead_dias", 15, "0"),
]

COLUNAS_SLA = [
    ("Projeto", "projeto", 30, None),
    ("Etapa", "etapa", 30, None),
    ("Responsável", "responsavel", 22, None),
    ("Data prevista", "prevista", 13, FORMATO_DATA),
    ("Concluída em", "concluida_em", 18, FORMATO_DATA_HORA),
    ("Atraso (dias)", "atraso", 12, "0"),
    ("Intervalo desde anterior", "cadencia", 22, FORMATO_DURACAO),
    ("Ciclo da etapa", "ciclo", 16, FORMATO_DURACAO),
]

COLUNAS_TEMPO = [
    ("Projeto", "projeto", 30, None),
    ("Membro", "membro", 22, None),
    ("Início", "inicio", 18, FORMATO_DATA_HORA),
    ("Fim", "fim", 18, FORMATO_DATA_HORA),
    ("Duração no período", "duracao", 18, FORMATO_DURACAO),
]

COLUNAS_MEMBROS = [
    ("Membro", "membro", 26, None),
    ("Tempo no período", "tempo", 18, FORMATO_DURACAO),
]

COLUNAS_EVENTOS = [
    ("Data", "quando", 18, FORMATO_DATA_HORA),
    ("Projeto", "projeto", 30, None),
    ("Tipo", "tipo", 20, None),
    ("Descrição", "descricao", 55, None),
    ("Autor", "autor", 22, None),
]


def _valor_planilha(valor):
    if isinstance(valor, datetime) and valor.tzinfo is not None:
        return valor.replace(tzinfo=None)
    return valor


def _escrever_aba(planilha, titulo, colunas, linhas):
    aba = planilha.create_sheet(titulo)
    cabecalho = PatternFill("solid", fgColor="1F4E78")
    for indice, (rotulo, _chave, largura, _formato) in enumerate(colunas, start=1):
        celula = aba.cell(row=1, column=indice, value=rotulo)
        celula.font = Font(bold=True, color="FFFFFF")
        celula.fill = cabecalho
        celula.alignment = Alignment(vertical="center")
        aba.column_dimensions[get_column_letter(indice)].width = largura
    for linha_idx, linha in enumerate(linhas, start=2):
        for col_idx, (_rotulo, chave, _largura, formato) in enumerate(colunas, start=1):
            celula = aba.cell(
                row=linha_idx, column=col_idx, value=_valor_planilha(linha.get(chave))
            )
            if formato:
                celula.number_format = formato
    aba.freeze_panes = "A2"
    if linhas:
        ultima = get_column_letter(len(colunas))
        aba.auto_filter.ref = f"A1:{ultima}{len(linhas) + 1}"
    return aba


def gerar_xlsx(relatorio):
    planilha = Workbook()
    resumo_aba = planilha.active
    aba_resumo = _escrever_aba(
        planilha,
        "Resumo",
        [("Indicador", "rotulo", 30, None), ("Valor", "valor", 28, None)],
        [{"rotulo": rotulo, "valor": valor} for rotulo, valor in relatorio["resumo"]],
    )
    for linha in range(2, len(relatorio["resumo"]) + 2):
        valor = aba_resumo.cell(row=linha, column=2).value
        celula = aba_resumo.cell(row=linha, column=1)
        celula.font = Font(bold=True)

    _escrever_aba(planilha, "Projetos", COLUNAS_PROJETOS, relatorio["projetos"])
    _escrever_aba(planilha, "Checklists (SLA)", COLUNAS_SLA, relatorio["sla"])
    _escrever_aba(planilha, "Tempo", COLUNAS_TEMPO, relatorio["sessoes"])
    _escrever_aba(planilha, "Tempo por membro", COLUNAS_MEMBROS, relatorio["membros"])
    _escrever_aba(planilha, "Eventos", COLUNAS_EVENTOS, relatorio["eventos"])

    del planilha["Sheet"]
    planilha.active = 0

    buffer = io.BytesIO()
    planilha.save(buffer)
    buffer.seek(0)
    return buffer
